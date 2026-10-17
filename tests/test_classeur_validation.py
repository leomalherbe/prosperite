"""Tests hors ligne de scripts/classeur_validation.py (classeurs Excel du gel 2, méthode 33 v1.3, D41)."""

from __future__ import annotations

import csv
import importlib.util
import sys
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

RACINE = Path(__file__).resolve().parent.parent


def _charger(nom: str):
    spec = importlib.util.spec_from_file_location(nom, RACINE / "scripts" / f"{nom}.py")
    module = importlib.util.module_from_spec(spec)
    sys.modules[nom] = module
    spec.loader.exec_module(module)
    return module


cla = _charger("classeur_validation")
ver = _charger("verifier_fiches")


@pytest.mark.parametrize("ligne, attendu", [
    ({"recommandation": "retenir", "variante_principale": "oui", "socle_propose": "oui"}, "retenu_socle"),
    ({"recommandation": "retenir", "variante_principale": "oui", "socle_propose": "non", "incontournable": "oui"},
     "retenu"),
    ({"recommandation": "retenir", "variante_principale": "oui", "socle_propose": "", "incontournable": "oui"},
     "retenu_socle"),                                                     # repli D37 (champ absent)
    ({"recommandation": "retenir", "variante_principale": "oui", "incontournable": "non"}, "retenu"),
    ({"recommandation": "retenir_variante", "variante_principale": "non", "socle_propose": "non"}, "retenu"),
    ({"recommandation": "ecarter", "variante_principale": "oui", "socle_propose": "non"}, "ecarte"),
    ({"recommandation": ""}, "a_discuter"),
])
def test_decision_par_defaut(ligne: dict, attendu: str):
    assert cla.decision_par_defaut(ligne) == attendu


def _fiche(fid: str, sous_theme: str, **surcharges) -> dict:
    fiche = {"id": fid, "libelle": f"Indicateur {fid}", "sous_theme": sous_theme, "famille": fid,
             "variante_principale": "oui", "definition": "Définition.", "sens": "haut", "unite": "%",
             "source_organisme": "INSEE", "source_jeu": "BDM", "source_serie_id": "001", "premiere_annee": 2000,
             "derniere_annee": 2024, "nb_points_depuis_2000": 25, "frequence": "annuelle", "type": "resultat",
             "declaratif": "non", "composite": "non", "incontournable": "non", "socle_propose": "non",
             "recommandation": "retenir", "motif_recommandation": "Motif."}
    fiche.update(surcharges)
    return fiche


@pytest.fixture
def dossiers(tmp_path: Path, monkeypatch) -> Path:
    monkeypatch.setattr(cla, "DOSSIER_CSV", tmp_path / "validation")
    monkeypatch.setattr(cla, "DOSSIER_XLSX", tmp_path / "validation" / "classeurs")
    monkeypatch.setattr(cla, "RACINE", tmp_path)
    (tmp_path / "validation").mkdir()
    return tmp_path


def _ecrire_csv(chemin: Path, fiches: list[dict], colonnes: list[str]) -> None:
    with chemin.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=colonnes, lineterminator="\n")
        w.writeheader()
        for fiche in fiches:
            ligne = ver.ligne_tableau(fiche)
            w.writerow({c: ligne[c] for c in colonnes})


def test_export_colonnes_et_synthese_par_sous_theme(dossiers: Path):
    fiches = [_fiche("a", "revenus", socle_propose="oui", incontournable="oui"), _fiche("b", "revenus"),
              _fiche("c", "pauvreté", socle_propose="oui"),
              _fiche("d", "pauvreté", recommandation="ecarter")]
    _ecrire_csv(dossiers / "validation" / "validation_dimtest.csv", fiches, ver.COLONNES_TABLEAU)
    sortie = cla.exporter("dimtest", "Dimension test")
    wb = load_workbook(sortie)
    ws = wb["Validation"]
    entetes = [c.value for c in ws[1]]
    assert "Sous-thème" in entetes and "Socle proposé" in entetes
    j_id, j_def = entetes.index("Identifiant") + 1, entetes.index("Décision par défaut (si vous ne remplissez rien)") + 1
    defauts = {ws.cell(row=i, column=j_id).value: ws.cell(row=i, column=j_def).value for i in range(2, ws.max_row + 1)}
    assert defauts == {"a": "retenu_socle", "b": "retenu", "c": "retenu_socle", "d": "ecarte"}
    # tri par sous-thème : « pauvreté » avant « revenus »
    assert [ws.cell(row=i, column=j_id).value for i in range(2, 6)] == ["c", "d", "a", "b"]
    sy = wb["Synthèse"]
    textes = [sy.cell(row=i, column=1).value for i in range(1, sy.max_row + 1)]
    assert "Socle par sous-thème (D41)" in textes and "pauvreté" in textes and "revenus" in textes
    rang = textes.index("revenus") + 1
    assert sy.cell(row=rang, column=3).value.startswith("=COUNTIFS(")
    assert "retenu_socle" in sy.cell(row=rang, column=4).value
    assert any(isinstance(sy.cell(row=i, column=2).value, str) and "socle_propose" not in sy.cell(row=i, column=2).value
               and sy.cell(row=i, column=1).value == "Socle proposé par l'agent (D41)" for i in range(1, sy.max_row + 1))


def test_export_ancien_tableau_complete_depuis_les_fiches(dossiers: Path):
    """Tableau v1.2 (sans sous_theme ni socle_propose) : les champs sont lus dans les fiches YAML."""
    fiche = _fiche("a", "revenus", incontournable="oui")
    del fiche["socle_propose"]
    (dossiers / "indicateurs" / "dimtest").mkdir(parents=True)
    (dossiers / "indicateurs" / "dimtest" / "a.yaml").write_text(yaml.safe_dump(fiche, allow_unicode=True),
                                                                 encoding="utf-8")
    _ecrire_csv(dossiers / "validation" / "validation_dimtest.csv", [fiche], ver.COLONNES_TABLEAU_V12)
    lignes = cla.lire_csv("dimtest")
    assert lignes[0]["sous_theme"] == "revenus" and lignes[0]["socle_propose"] == ""
    assert cla.decision_par_defaut(lignes[0]) == "retenu_socle"


def test_import_reporte_les_decisions(dossiers: Path):
    fiches = [_fiche("a", "revenus", socle_propose="oui"), _fiche("b", "revenus")]
    chemin_csv = dossiers / "validation" / "validation_dimtest.csv"
    _ecrire_csv(chemin_csv, fiches, ver.COLONNES_TABLEAU)
    sortie = cla.exporter("dimtest")
    wb = load_workbook(sortie)
    ws = wb["Validation"]
    entetes = [c.value for c in ws[1]]
    j_id = entetes.index("Identifiant") + 1
    j_dec = entetes.index("Votre décision (liste déroulante)") + 1
    j_com = entetes.index("Votre commentaire") + 1
    for i in range(2, ws.max_row + 1):
        if ws.cell(row=i, column=j_id).value == "b":
            ws.cell(row=i, column=j_dec, value="retenu_socle")
            ws.cell(row=i, column=j_com, value="à cocher aussi")
    wb.save(sortie)
    assert cla.importer("dimtest") == 1
    relu = {r["id"]: r for r in csv.DictReader(chemin_csv.open(encoding="utf-8"))}
    assert relu["b"]["decision_leo"] == "retenu_socle" and relu["b"]["commentaire_leo"] == "à cocher aussi"
    assert relu["a"]["decision_leo"] == "" and list(relu["a"]) == ver.COLONNES_TABLEAU


def _saisir(sortie: Path, valeurs: dict[str, tuple[str, str]]) -> None:
    wb = load_workbook(sortie)
    ws = wb["Validation"]
    entetes = [c.value for c in ws[1]]
    j_id = entetes.index("Identifiant") + 1
    j_dec = entetes.index("Votre décision (liste déroulante)") + 1
    j_com = entetes.index("Votre commentaire") + 1
    for i in range(2, ws.max_row + 1):
        ident = ws.cell(row=i, column=j_id).value
        if ident in valeurs:
            ws.cell(row=i, column=j_dec, value=valeurs[ident][0] or None)
            ws.cell(row=i, column=j_com, value=valeurs[ident][1] or None)
    wb.save(sortie)


def test_reexport_conserve_les_decisions_non_importees(dossiers: Path):
    """Régénérer le classeur ne perd pas une décision saisie et pas encore importée ; le CSV prime en cas de
    désaccord."""
    fiches = [_fiche("a", "revenus", socle_propose="oui"), _fiche("b", "revenus"), _fiche("c", "pauvreté")]
    chemin_csv = dossiers / "validation" / "validation_dimtest.csv"
    _ecrire_csv(chemin_csv, fiches, ver.COLONNES_TABLEAU)
    sortie = cla.exporter("dimtest")
    _saisir(sortie, {"b": ("ecarte", "doublon"), "c": ("a_discuter", "")})
    # le CSV a reçu entre-temps une autre décision pour c
    lignes = list(csv.DictReader(chemin_csv.open(encoding="utf-8", newline="")))
    for l in lignes:
        if l["id"] == "c":
            l["decision_leo"] = "retenu"
    with chemin_csv.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=ver.COLONNES_TABLEAU, lineterminator="\n")
        w.writeheader()
        w.writerows(lignes)
    cla.exporter("dimtest")
    relues = cla.lire_decisions_classeur(sortie)
    assert relues["b"] == ("ecarte", "doublon")          # reprise du classeur existant
    assert relues["c"] == ("retenu", "")                  # le CSV prime
    assert relues["a"] == ("", "")


def test_import_ne_vide_jamais_une_decision_et_ecrit_en_lf(dossiers: Path):
    fiches = [_fiche("a", "revenus", socle_propose="oui"), _fiche("b", "revenus")]
    chemin_csv = dossiers / "validation" / "validation_dimtest.csv"
    _ecrire_csv(chemin_csv, fiches, ver.COLONNES_TABLEAU)
    sortie = cla.exporter("dimtest")
    _saisir(sortie, {"a": ("retenu", "à garder")})
    assert cla.importer("dimtest") == 1
    # classeur où la décision de a a été effacée et un commentaire ajouté à b
    _saisir(sortie, {"a": ("", ""), "b": ("", "note seule")})
    assert cla.importer("dimtest") == 1
    relu = {r["id"]: r for r in csv.DictReader(chemin_csv.open(encoding="utf-8", newline=""))}
    assert (relu["a"]["decision_leo"], relu["a"]["commentaire_leo"]) == ("retenu", "à garder")
    assert (relu["b"]["decision_leo"], relu["b"]["commentaire_leo"]) == ("", "note seule")
    assert b"\r\n" not in chemin_csv.read_bytes()
    assert list(relu["a"]) == ver.COLONNES_TABLEAU
