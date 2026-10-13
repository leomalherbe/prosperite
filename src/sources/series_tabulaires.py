"""Inventaire des séries contenues dans un fichier de données (mission M04, phase 2).

Fonctions génériques, sans accès réseau, utilisées par `scripts/collecte_sources.py`
(exécuté par GitHub Actions) pour transformer un fichier téléchargé (CSV, XLSX,
ZIP, JSON) en une **liste de séries** : une ligne par série avec sa clé, son
libellé, sa première et sa dernière période, le nombre de périodes renseignées,
le nombre de valeurs distinctes et la dernière valeur.

Le but est de **repérer et vérifier** les séries disponibles (indicateurs
candidats), pas de produire des données pour l'indice : aucune valeur n'est
recalculée ni corrigée, la dernière valeur n'est qu'un témoin de lecture.

Modes (chaîne « nom;cle=valeur;… ») :
- ``series``     : table « longue » (une ligne par observation). Paramètres :
                   ``cles`` (colonnes séparées par « | », « - » = aucune, défaut : détection),
                   ``temps`` (colonne de période, défaut : détection),
                   ``valeurs`` (colonnes numériques, « * » = détection),
                   ``filtre`` (« Col=Val,Col2~motif »), ``libelle``, ``unite``,
                   ``agg=somme`` (somme des lignes d'une même clé et période),
                   ``feuille``, ``entete`` (ligne d'en-tête d'un classeur, 0 = 1re),
                   ``zip`` (membre d'une archive), ``sep``, ``encodage``.
- ``large``      : table « large » (une colonne par année). Paramètres : ``cles``,
                   ``libelle``, ``unite``, ``filtre``, ``motif_colonne`` (regex à
                   deux groupes variable/année ou un groupe année ; défaut ``^A(\\d{4})$``).
- ``couverture`` : microdonnées d'enquête ; pour chaque variable, années où elle
                   est renseignée. Paramètre ``annee`` (colonne de l'année).
- ``lignes``     : lignes (CSV) ou rangées (XLSX, y compris dans un ZIP) qui
                   contiennent le motif ``motif`` (regex), avec les premières lignes.
- ``classeur``   : description d'un classeur (``lignes`` par feuille, ``feuilles``).
- ``ods_champs`` : catalogue Opendatasoft exporté en JSON → jeux et champs.
- ``dido``       : catalogue de la plateforme DiDo du SDES → jeux et fichiers.
- ``agregat``    : somme de colonnes par clé (gros CSV lu par morceaux) ;
                   paramètres ``cle``, ``sommes``, ``ratio=num/den``.
- ``bm``         : réponse JSON de l'API de la Banque mondiale.
"""

from __future__ import annotations

import io
import json
import re
import zipfile
from pathlib import Path

COLONNES_SERIES = [
    "source", "requete", "serie_id", "libelle", "unite", "details",
    "premiere_periode", "derniere_periode", "n_periodes", "n_valeurs_distinctes",
    "derniere_valeur", "periodes",
]
NOMS_TEMPS = ("time_period", "annee", "année", "annees", "year", "periode", "période",
              "unite_temps", "an", "aaaa", "an_mesure", "annee_de_l_enquete", "date", "millesime", "millésime",
              "session", "rentree_scolaire")
MAX_SERIES = 20_000


# --------------------------------------------------------------------------- outils


def parametres(mode: str) -> tuple[str, dict]:
    """« series;cles=A|B;temps=T » → ("series", {"cles": "A|B", "temps": "T"})."""
    morceaux = [m for m in mode.split(";") if m != ""]
    if not morceaux:
        return "", {}
    nom, p = morceaux[0].strip(), {}
    for m in morceaux[1:]:
        if "=" in m:
            k, v = m.split("=", 1)
            p[k.strip()] = v
    return nom, p


def detecter_encodage(echantillon: bytes) -> str:
    """utf-8 si l'échantillon se décode (une coupure en fin d'échantillon est tolérée), sinon cp1252."""
    try:
        echantillon.decode("utf-8")
        return "utf-8-sig"
    except UnicodeDecodeError as exc:
        if exc.start >= len(echantillon) - 3:
            return "utf-8-sig"
        return "cp1252"


def detecter_separateur(premiere_ligne: str) -> str:
    return max([",", ";", "\t", "|"], key=premiere_ligne.count)


def compresser_periodes(periodes: list[str]) -> str:
    """['2000','2001','2002','2005'] → '2000-2002,2005' (années) ; autres périodes listées telles quelles."""
    if not periodes:
        return ""
    if all(re.fullmatch(r"\d{4}", p) for p in periodes):
        ans = sorted({int(p) for p in periodes})
        plages, debut, prec = [], ans[0], ans[0]
        for a in ans[1:]:
            if a == prec + 1:
                prec = a
                continue
            plages.append(f"{debut}-{prec}" if prec > debut else f"{debut}")
            debut = prec = a
        plages.append(f"{debut}-{prec}" if prec > debut else f"{debut}")
        return ",".join(plages)
    if len(periodes) > 12:
        return ",".join(periodes[:6]) + ",…," + ",".join(periodes[-6:])
    return ",".join(periodes)


def en_nombre(serie):
    import pandas as pd

    s = serie.astype(str).str.strip().str.replace(" ", "", regex=False).str.replace("\xa0", "", regex=False)
    s = s.str.replace(r"^(-?\d+),(\d+)$", r"\1.\2", regex=True)
    return pd.to_numeric(s, errors="coerce")


def colonne(df, nom: str) -> str:
    """Nom de colonne réel, sans tenir compte de la casse ni des espaces."""
    if nom in df.columns:
        return nom
    cible = nom.strip().lower()
    for c in df.columns:
        if str(c).strip().lower() == cible:
            return c
    raise KeyError(f"colonne « {nom} » absente ; colonnes : {list(df.columns)[:60]}")


def appliquer_filtre(df, filtre: str):
    """« code=FRA,Zone~^France » : égalité (=) ou motif (~), conditions cumulées."""
    if not filtre:
        return df
    for cond in filtre.split(","):
        if "~" in cond:
            c, motif = cond.split("~", 1)
            df = df[df[colonne(df, c)].astype(str).str.contains(motif, regex=True, na=False)]
        elif "=" in cond:
            c, val = cond.split("=", 1)
            df = df[df[colonne(df, c)].astype(str).str.strip() == val]
    return df


# --------------------------------------------------------------------------- lecture


def membre_zip(contenu: bytes, membre: str | None) -> tuple[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(contenu)) as z:
        noms = [n for n in z.namelist() if not n.endswith("/")]
        if membre:
            choix = [n for n in noms if n == membre or n.endswith(membre)]
        else:
            choix = [n for n in noms if n.lower().endswith((".csv", ".xlsx", ".xls", ".txt"))]
        if not choix:
            raise ValueError(f"membre introuvable dans l'archive ; membres : {noms[:50]}")
        return choix[0], z.read(choix[0])


def lire_classeur_df(contenu: bytes, feuille: str | None = None, entete: int = 0):
    import openpyxl
    import pandas as pd

    wb = openpyxl.load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    if not wb.sheetnames:
        wb = openpyxl.load_workbook(io.BytesIO(contenu), read_only=False, data_only=True)
    ws = wb[feuille] if feuille else wb[wb.sheetnames[0]]
    lignes = list(ws.iter_rows(values_only=True))
    if len(lignes) <= entete:
        return pd.DataFrame()
    noms = [("" if v is None else str(v).strip()) or f"col{i}" for i, v in enumerate(lignes[entete])]
    corps = [["" if v is None else str(v) for v in l] for l in lignes[entete + 1:]]
    largeur = len(noms)
    corps = [(l + [""] * largeur)[:largeur] for l in corps]
    return pd.DataFrame(corps, columns=noms)


def lire_table(chemin: Path, p: dict):
    """Charge un CSV, un XLSX ou un membre de ZIP en DataFrame de chaînes."""
    import pandas as pd

    contenu = Path(chemin).read_bytes()
    if contenu[:4] == b"PK\x03\x04":
        # XLSX ou ZIP : un XLSX est lui-même une archive contenant xl/workbook.xml
        with zipfile.ZipFile(io.BytesIO(contenu)) as z:
            est_classeur = "xl/workbook.xml" in z.namelist()
        if not est_classeur:
            _, contenu = membre_zip(contenu, p.get("zip"))
        if contenu[:4] == b"PK\x03\x04":
            return lire_classeur_df(contenu, p.get("feuille"), int(p.get("entete", 0)))
    encodage = p.get("encodage") or detecter_encodage(contenu[:2_000_000])
    texte = contenu.decode(encodage, errors="replace")
    sep = p.get("sep") or detecter_separateur(texte.split("\n", 1)[0])
    return pd.read_csv(io.StringIO(texte), sep=sep, dtype=str, low_memory=False, on_bad_lines="skip",
                       keep_default_na=False, na_values=[""])


# --------------------------------------------------------------------------- séries


def libelles_sdmx(df) -> dict:
    """CSV SDMX « avec libellés » : chaque colonne de code est suivie de sa colonne de libellé."""
    if "STRUCTURE" not in df.columns:
        return {}
    cols = list(df.columns)
    return {c: cols[i + 1] for i, c in enumerate(cols[:-1])
            if re.fullmatch(r"[A-Z0-9_]+", str(c)) and not re.fullmatch(r"[A-Z0-9_]+", str(cols[i + 1]))}


def detecter_temps(df) -> str | None:
    for nom in NOMS_TEMPS:
        for c in df.columns:
            if str(c).strip().lower() == nom:
                return c
    for c in df.columns:
        echant = df[c].dropna().astype(str).head(200)
        if len(echant) and echant.str.fullmatch(r"(19|20)\d{2}([-/ ]?(Q[1-4]|S[12]|T[1-4]|\d{2}))?").mean() > 0.95:
            return c
    return None


def resumer_series(df, p: dict, source: str = "", requete: str = "") -> list[dict]:
    """Une ligne par (clé, colonne de valeur) : couverture temporelle et dernière valeur."""
    import pandas as pd

    df = appliquer_filtre(df, p.get("filtre", ""))
    if df.empty:
        return []
    etiquettes = libelles_sdmx(df)
    temps = colonne(df, p["temps"]) if p.get("temps") else detecter_temps(df)
    if temps is None:
        raise ValueError(f"colonne de période non trouvée ; colonnes : {list(df.columns)[:60]}")
    libelle_col = colonne(df, p["libelle"]) if p.get("libelle") else None
    unite_col = colonne(df, p["unite"]) if p.get("unite") else None
    exclues = {temps, libelle_col, unite_col} | set(etiquettes.values())
    if p.get("valeurs") and p["valeurs"] != "*":
        valeurs = [colonne(df, v) for v in p["valeurs"].split("|")]
    else:
        valeurs = []
        for c in df.columns:
            if c in exclues:
                continue
            s = df[c].dropna()
            s = s[s.astype(str).str.strip() != ""]
            if len(s) and en_nombre(s).notna().mean() >= 0.8:
                valeurs.append(c)
    if p.get("cles") == "-":  # aucune clé : une seule série par colonne de valeur
        cles = []
    elif p.get("cles"):
        cles = [colonne(df, c) for c in p["cles"].split("|")]
    else:
        cles = [c for c in df.columns
                if c not in exclues and c not in valeurs and 1 <= df[c].nunique(dropna=True) <= 300]
        cles = [c for c in cles if c not in ("STRUCTURE", "STRUCTURE_ID", "STRUCTURE_NAME", "ACTION")]
    if unite_col is None and "UNIT_MEASURE" in df.columns:
        unite_col = etiquettes.get("UNIT_MEASURE", "UNIT_MEASURE")

    lignes = []
    travail = df.copy()
    travail["__t"] = travail[temps].astype(str).str.strip()
    for v in valeurs:
        travail["__v"] = en_nombre(travail[v])
        sous = travail[travail["__v"].notna() & (travail["__t"] != "")]
        if sous.empty:
            continue
        groupes = sous.groupby(cles, dropna=False, sort=True) if cles else [((), sous)]
        for cle, g in groupes:
            cle = cle if isinstance(cle, tuple) else (cle,)
            if p.get("agg") == "somme":
                par_t = g.groupby("__t")["__v"].sum()
            else:
                par_t = g.groupby("__t")["__v"].last()
            par_t = par_t.sort_index()
            periodes = [str(t) for t in par_t.index]
            id_ = ".".join("" if (isinstance(x, float) and x != x) else str(x) for x in cle)
            if len(valeurs) > 1:
                id_ = f"{id_}#{v}" if id_ else str(v)
            premier = g.iloc[0]
            if libelle_col:
                lib = str(premier[libelle_col])
            elif cles and cles[0] in etiquettes:
                lib = str(premier[etiquettes[cles[0]]])
            else:
                lib = str(v)
            details = " ; ".join(
                f"{c}={premier[etiquettes[c]] if c in etiquettes else premier[c]}" for c in cles[1:])
            lignes.append({
                "source": source, "requete": requete, "serie_id": id_, "libelle": lib,
                "unite": "" if unite_col is None else str(premier.get(unite_col, "")),
                "details": details,
                "premiere_periode": periodes[0], "derniere_periode": periodes[-1],
                "n_periodes": len(periodes), "n_valeurs_distinctes": int(par_t.round(6).nunique()),
                "derniere_valeur": f"{par_t.iloc[-1]:.6g}", "periodes": compresser_periodes(periodes),
            })
            if len(lignes) >= MAX_SERIES:
                return lignes
    return lignes


def resumer_large(df, p: dict, source: str = "", requete: str = "") -> list[dict]:
    """Table large : colonnes annuelles (A2000… ou hdi_2000…) ; une ligne par série."""
    df = appliquer_filtre(df, p.get("filtre", ""))
    motif = re.compile(p.get("motif_colonne") or r"^A(\d{4})$")
    colonnes = {}
    for c in df.columns:
        m = motif.match(str(c))
        if m:
            if m.lastindex and m.lastindex >= 2:
                colonnes[c] = (m.group(1), m.group(2))
            else:
                colonnes[c] = ("", m.group(1))
    if not colonnes:
        raise ValueError(f"aucune colonne ne correspond à {motif.pattern} ; colonnes : {list(df.columns)[:60]}")
    cles = [colonne(df, c) for c in p.get("cles", "").split("|") if c]
    libelle_col = colonne(df, p["libelle"]) if p.get("libelle") else None
    unite_col = colonne(df, p["unite"]) if p.get("unite") else None
    numeriques = {c: en_nombre(df[c]) for c in colonnes}
    lignes = []
    for pos, (_, r) in enumerate(df.iterrows()):
        par_variable: dict[str, dict[str, float]] = {}
        for c, (var, an) in colonnes.items():
            x = numeriques[c].iloc[pos]
            if x == x:  # non NaN
                par_variable.setdefault(var, {})[an] = float(x)
        for var, obs in par_variable.items():
            ans = sorted(obs)
            id_ = ".".join(str(r[c]) for c in cles)
            if var:
                id_ = f"{id_}#{var}" if id_ else var
            lignes.append({
                "source": source, "requete": requete, "serie_id": id_,
                "libelle": str(r[libelle_col]) if libelle_col else var,
                "unite": str(r[unite_col]) if unite_col else "", "details": "",
                "premiere_periode": ans[0], "derniere_periode": ans[-1], "n_periodes": len(ans),
                "n_valeurs_distinctes": len({round(v, 6) for v in obs.values()}),
                "derniere_valeur": f"{obs[ans[-1]]:.6g}", "periodes": compresser_periodes(ans),
            })
    return lignes


def couverture_microdonnees(chemin: Path, p: dict, taille_morceau: int = 50_000) -> list[dict]:
    """Pour chaque variable d'un fichier de microdonnées : années où elle est renseignée."""
    import pandas as pd

    with open(chemin, "rb") as f:
        echantillon = f.read(2_000_000)
    encodage = p.get("encodage") or detecter_encodage(echantillon)
    premiere = echantillon.decode(encodage, errors="replace").split("\n", 1)[0]
    sep = p.get("sep") or detecter_separateur(premiere)
    col_annee = p.get("annee", "annee")
    comptes: dict[str, dict[str, int]] = {}
    effectifs: dict[str, int] = {}
    for morceau in pd.read_csv(chemin, sep=sep, dtype=str, encoding=encodage, chunksize=taille_morceau,
                               low_memory=False, on_bad_lines="skip", keep_default_na=False, na_values=[""]):
        annee = colonne(morceau, col_annee)
        renseigne = morceau.notna() & (morceau.astype(str).apply(lambda s: s.str.strip()) != "")
        renseigne[annee] = True
        somme = renseigne.groupby(morceau[annee].astype(str).str.strip()).sum()
        for a, n in morceau[annee].astype(str).str.strip().value_counts().items():
            effectifs[a] = effectifs.get(a, 0) + int(n)
        for var in somme.columns:
            d = comptes.setdefault(var, {})
            for a, n in somme[var].items():
                if n:
                    d[a] = d.get(a, 0) + int(n)
    lignes = []
    for var, d in comptes.items():
        ans = sorted(d)
        if not ans:
            continue
        der = ans[-1]
        lignes.append({
            "variable": var, "n_annees": len(ans), "premiere_annee": ans[0], "derniere_annee": der,
            "annees": compresser_periodes(ans),
            "taux_reponse_derniere_annee": f"{d[der] / max(effectifs.get(der, 1), 1):.3f}",
        })
    lignes.append({"variable": "__effectifs__", "n_annees": len(effectifs), "premiere_annee": min(effectifs, default=""),
                   "derniere_annee": max(effectifs, default=""),
                   "annees": " ".join(f"{a}:{n}" for a, n in sorted(effectifs.items())),
                   "taux_reponse_derniere_annee": ""})
    return lignes


def colonne_correspondante(contenu: bytes, motif: str, entete: int = 10) -> str:
    """Classeur « transposé » (un pays par colonne) : pour chaque feuille, colonne dont l'en-tête
    (dans les `entete` premières lignes) correspond au motif, avec le libellé de la première colonne."""
    import openpyxl

    rx = re.compile(motif)
    wb = openpyxl.load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    sortie = [f"Feuilles ({len(wb.sheetnames)}) : {wb.sheetnames}"]
    for f in wb.sheetnames:
        ws = wb[f]
        if not hasattr(ws, "iter_rows"):
            continue
        rangees = list(ws.iter_rows(values_only=True))
        j = None
        for rangee in rangees[:entete]:
            for k, v in enumerate(rangee):
                if v is not None and rx.search(str(v)):
                    j = k
                    break
            if j is not None:
                break
        if j is None:
            sortie.append(f"\n=== Feuille « {f} » : motif absent des {entete} premières lignes ===")
            continue
        sortie.append(f"\n=== Feuille « {f} » : colonne {j} ===")
        for i, rangee in enumerate(rangees):
            lib = "" if not rangee or rangee[0] is None else str(rangee[0])
            val = "" if j >= len(rangee) or rangee[j] is None else str(rangee[j])
            if lib or val:
                sortie.append(f"[{i}] {lib[:200]} | {val[:60]}")
    return "\n".join(sortie)


def distribution_ponderee(chemin: Path, p: dict, taille_morceau: int = 50_000) -> list[dict]:
    """Microdonnées pondérées : part pondérée (%) de chaque modalité, par année, pour les variables
    dont le nom correspond au motif `variables` (regex). Les valeurs vides sont exclues du dénominateur."""
    import pandas as pd

    with open(chemin, "rb") as f:
        echantillon = f.read(2_000_000)
    encodage = p.get("encodage") or detecter_encodage(echantillon)
    sep = p.get("sep") or detecter_separateur(echantillon.decode(encodage, errors="replace").split("\n", 1)[0])
    col_annee, col_poids = p.get("annee", "annee"), p.get("poids", "poids")
    rx = re.compile(p.get("variables", ".*"))
    sommes: dict[tuple, float] = {}
    for morceau in pd.read_csv(chemin, sep=sep, dtype=str, encoding=encodage, chunksize=taille_morceau,
                               low_memory=False, on_bad_lines="skip", keep_default_na=False, na_values=[""]):
        annee = morceau[colonne(morceau, col_annee)].astype(str).str.strip()
        poids = en_nombre(morceau[colonne(morceau, col_poids)]).fillna(0)
        for var in morceau.columns:
            if var in (col_annee, col_poids) or not rx.fullmatch(str(var)):
                continue
            valeurs = morceau[var].astype(str).str.strip()
            ok = morceau[var].notna() & (valeurs != "")
            if not ok.any():
                continue
            g = poids[ok].groupby([annee[ok], valeurs[ok]]).sum()
            for (a, m), w in g.items():
                sommes[(var, a, m)] = sommes.get((var, a, m), 0.0) + float(w)
    totaux: dict[tuple, float] = {}
    for (var, a, m), w in sommes.items():
        totaux[(var, a)] = totaux.get((var, a), 0.0) + w
    return [{"variable": var, "annee": a, "modalite": m, "part_ponderee": f"{100 * w / totaux[(var, a)]:.2f}"}
            for (var, a, m), w in sorted(sommes.items()) if totaux[(var, a)]]


def lignes_correspondantes(contenu: bytes, motif: str, nom: str = "", max_lignes: int = 400,
                           entete: int = 12) -> str:
    """Premières lignes et lignes contenant le motif, pour un CSV, un XLSX ou un ZIP."""
    rx = re.compile(motif)
    if contenu[:4] == b"PK\x03\x04":
        with zipfile.ZipFile(io.BytesIO(contenu)) as z:
            noms = z.namelist()
            if "xl/workbook.xml" not in noms:
                sorties = []
                for n in noms:
                    if n.lower().endswith((".csv", ".xlsx", ".txt")):
                        sorties.append(f"##### Membre {n}\n" + lignes_correspondantes(z.read(n), motif, n, max_lignes, entete))
                return "\n".join(sorties) or f"Archive sans CSV/XLSX : {noms[:50]}"
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
        if not wb.sheetnames:
            wb = openpyxl.load_workbook(io.BytesIO(contenu), read_only=False, data_only=True)
        sortie = [f"Feuilles ({len(wb.sheetnames)}) : {wb.sheetnames}"]
        for f in wb.sheetnames:
            ws = wb[f]
            if not hasattr(ws, "iter_rows"):
                continue
            sortie.append(f"\n=== Feuille « {f} » ===")
            trouves = 0
            for i, rangee in enumerate(ws.iter_rows(values_only=True)):
                cellules = ["" if v is None else str(v) for v in rangee]
                while cellules and cellules[-1] == "":
                    cellules.pop()
                texte = " | ".join(cellules)
                if i < entete:
                    sortie.append(f"[{i}] {texte[:3000]}")
                elif any(rx.search(c) for c in cellules):
                    sortie.append(f"[{i}] {texte[:6000]}")
                    trouves += 1
                    if trouves >= max_lignes:
                        sortie.append("… (tronqué)")
                        break
        return "\n".join(sortie)
    encodage = detecter_encodage(contenu[:2_000_000])
    lignes = contenu.decode(encodage, errors="replace").splitlines()
    sortie = [f"Encodage détecté : {encodage} ; {len(lignes)} lignes"]
    sortie += [f"[{i}] {l[:3000]}" for i, l in enumerate(lignes[:min(entete, 3)])]
    trouves = 0
    for i, l in enumerate(lignes[3:], 3):
        if rx.search(l):
            sortie.append(f"[{i}] {l[:6000]}")
            trouves += 1
            if trouves >= max_lignes:
                sortie.append("… (tronqué)")
                break
    return "\n".join(sortie)


# --------------------------------------------------------------------------- catalogues


def catalogue_ods(contenu: bytes) -> list[dict]:
    """Export JSON du catalogue Opendatasoft (…/catalog/exports/json) → jeux et champs."""
    donnees = json.loads(contenu.decode("utf-8", errors="replace"))
    if isinstance(donnees, dict):
        donnees = donnees.get("results") or donnees.get("datasets") or []
    lignes = []
    for d in donnees:
        metas = (d.get("metas") or {}).get("default", {}) if isinstance(d.get("metas"), dict) else {}
        dcat = (d.get("metas") or {}).get("dcat", {}) if isinstance(d.get("metas"), dict) else {}
        champs = d.get("fields") or []
        lignes.append({
            "dataset_id": d.get("dataset_id") or d.get("datasetid", ""),
            "titre": metas.get("title", ""),
            "enregistrements": metas.get("records_count", ""),
            "modifie": metas.get("modified", ""),
            "couverture": f"{dcat.get('temporal_coverage_start', '') or ''}/{dcat.get('temporal_coverage_end', '') or ''}".strip("/"),
            "theme": ";".join(metas.get("theme") or []) if isinstance(metas.get("theme"), list) else metas.get("theme", ""),
            "pieces_jointes": len(d.get("attachments") or []),
            "champs": " ; ".join(f"{c.get('name')}({c.get('type')})" for c in champs),
        })
    return lignes


def catalogue_dido(contenu: bytes) -> list[dict]:
    """Réponse JSON de …/dido/api/v1/datasets → une ligne par fichier de données (rid)."""
    donnees = json.loads(contenu.decode("utf-8", errors="replace"))
    jeux = donnees.get("data", donnees) if isinstance(donnees, dict) else donnees
    lignes = []
    for j in jeux:
        tc = j.get("temporal_coverage") or {}
        base = {"dataset_id": j.get("id", ""), "jeu": j.get("title", ""), "theme": j.get("topic", ""),
                "granularite": (j.get("spatial") or {}).get("granularity", ""),
                "couverture_jeu": f"{tc.get('start', '')}/{tc.get('end', '')}", "frequence": j.get("frequency", "")}
        fichiers = j.get("datafiles") or []
        if not fichiers:
            lignes.append({**base, "rid": "", "fichier": "", "couverture_fichier": ""})
        for f in fichiers:
            tcf = f.get("temporal_coverage") or {}
            lignes.append({**base, "rid": f.get("rid", ""), "fichier": f.get("title", ""),
                           "couverture_fichier": f"{tcf.get('start', '')}/{tcf.get('end', '')}"})
    return lignes


def series_banque_mondiale(contenu: bytes, source: str = "", requete: str = "") -> list[dict]:
    """Réponse JSON [métadonnées, observations] de l'API v2 de la Banque mondiale."""
    donnees = json.loads(contenu.decode("utf-8", errors="replace"))
    if not isinstance(donnees, list) or len(donnees) < 2 or not donnees[1]:
        raise ValueError(f"réponse sans observations : {str(donnees)[:300]}")
    obs = {}
    lib = code = ""
    for o in donnees[1]:
        lib = (o.get("indicator") or {}).get("value", lib)
        code = (o.get("indicator") or {}).get("id", code)
        if o.get("value") is not None:
            obs[str(o.get("date"))] = float(o["value"])
    ans = sorted(obs)
    if not ans:
        raise ValueError("aucune valeur renseignée")
    return [{"source": source, "requete": requete, "serie_id": code, "libelle": lib, "unite": "", "details": "",
             "premiere_periode": ans[0], "derniere_periode": ans[-1], "n_periodes": len(ans),
             "n_valeurs_distinctes": len({round(v, 6) for v in obs.values()}),
             "derniere_valeur": f"{obs[ans[-1]]:.6g}", "periodes": compresser_periodes(ans)}]


def agreger(chemin: Path, p: dict, taille_morceau: int = 500_000) -> list[dict]:
    """Somme de colonnes par clé sur un gros CSV lu par morceaux (ex. résultats électoraux)."""
    import pandas as pd

    with open(chemin, "rb") as f:
        echantillon = f.read(2_000_000)
    encodage = p.get("encodage") or detecter_encodage(echantillon)
    sep = p.get("sep") or detecter_separateur(echantillon.decode(encodage, errors="replace").split("\n", 1)[0])
    cle = p["cle"]
    sommes = p["sommes"].split("|")
    total = None
    for morceau in pd.read_csv(chemin, sep=sep, dtype=str, encoding=encodage, usecols=[cle] + sommes,
                               chunksize=taille_morceau, on_bad_lines="skip"):
        for c in sommes:
            morceau[c] = en_nombre(morceau[c])
        s = morceau.groupby(cle)[sommes].sum()
        n = morceau.groupby(cle).size().rename("n_lignes")
        s = s.join(n)
        total = s if total is None else total.add(s, fill_value=0)
    lignes = []
    if total is None:
        return lignes
    for k, r in total.sort_index().iterrows():
        ligne = {cle: k, "n_lignes": int(r["n_lignes"])}
        for c in sommes:
            ligne[c] = int(r[c]) if r[c] == r[c] else ""
        if p.get("ratio"):
            num, den = p["ratio"].split("/")
            ligne[f"ratio_{num}_{den}"] = f"{100 * r[num] / r[den]:.2f}" if r[den] else ""
        lignes.append(ligne)
    return lignes
