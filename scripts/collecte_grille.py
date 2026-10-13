"""Mission M03 (phase 2) — contre-vérification du Doughnut 3.0 : documents supplémentaires.

Script **ponctuel** exécuté par le robot « Collecte » (le bac à sable des sessions ne peut
pas télécharger l'archive). Il récupère l'archive des documents supplémentaires de
Fanning & Raworth (2025, Nature 646, 47-56), déposée sur ORA (Oxford University Research
Archive, licence CC BY 4.0), et en extrait le texte (PDF, DOCX) et le début des feuilles
(XLSX, CSV) dans `journaux/m03_doughnut_supplementaire/`, pour lecture par l'agent M03.

Aucune donnée n'est écrite dans `catalogues/` ni `donnees/`. Le script est retiré de la
branche après exécution (il n'a pas vocation à tourner chaque mois sur `main`).
"""

from __future__ import annotations

import io
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import requests

RACINE = Path(__file__).resolve().parent.parent
SORTIE = RACINE / "journaux" / "m03_doughnut_supplementaire"
URL_ORA_ZIP = "https://ora.ox.ac.uk/objects/uuid:cb7f2af6-96de-4744-9b44-185ce77c6b14/files/r6h440v46z"
ENTETES = {"User-Agent": "observatoire-prosperite/0.1 (projet de recherche ; github.com/leomalherbe/prosperite)"}
MAX_LIGNES_FEUILLE = 400


def _pypdf():
    try:
        import pypdf  # noqa: F401
    except ImportError:
        subprocess.run([sys.executable, "-m", "pip", "install", "-q", "pypdf>=4"], check=True)
    import pypdf

    return pypdf


def _nom_sur(nom: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", nom)[:120]


def extraire_pdf(contenu: bytes) -> str:
    pypdf = _pypdf()
    lecteur = pypdf.PdfReader(io.BytesIO(contenu))
    morceaux = []
    for i, page in enumerate(lecteur.pages, start=1):
        morceaux.append(f"\n\n=== page {i} ===\n")
        morceaux.append(page.extract_text() or "")
    return "".join(morceaux)


def extraire_docx(contenu: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(contenu)) as z:
        xml = z.read("word/document.xml").decode("utf-8", errors="replace")
    xml = re.sub(r"</w:p>", "\n", xml)
    xml = re.sub(r"</w:tc>", " | ", xml)
    return re.sub(r"<[^>]+>", "", xml)


def extraire_xlsx(contenu: bytes) -> str:
    import openpyxl

    classeur = openpyxl.load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    morceaux = []
    for feuille in classeur.worksheets:
        morceaux.append(f"\n\n=== feuille « {feuille.title} » (dimensions déclarées : {feuille.max_row} lignes x {feuille.max_column} colonnes) ===\n")
        for n, ligne in enumerate(feuille.iter_rows(values_only=True), start=1):
            if n > MAX_LIGNES_FEUILLE:
                morceaux.append(f"[… tronqué après {MAX_LIGNES_FEUILLE} lignes]\n")
                break
            morceaux.append(" | ".join("" if v is None else str(v) for v in ligne) + "\n")
    return "".join(morceaux)


def traiter(nom: str, contenu: bytes) -> str | None:
    bas = nom.lower()
    if bas.endswith(".pdf"):
        return extraire_pdf(contenu)
    if bas.endswith(".docx"):
        return extraire_docx(contenu)
    if bas.endswith((".xlsx", ".xlsm")):
        return extraire_xlsx(contenu)
    if bas.endswith((".csv", ".txt", ".md", ".r", ".json")):
        texte = contenu.decode("utf-8", errors="replace")
        lignes = texte.splitlines()
        if len(lignes) > MAX_LIGNES_FEUILLE:
            texte = "\n".join(lignes[:MAX_LIGNES_FEUILLE]) + f"\n[… tronqué après {MAX_LIGNES_FEUILLE} lignes sur {len(lignes)}]"
        return texte
    return None


def parcourir_zip(contenu: bytes, prefixe: str, inventaire: list[str]) -> None:
    with zipfile.ZipFile(io.BytesIO(contenu)) as z:
        for info in z.infolist():
            if info.is_dir():
                continue
            nom = f"{prefixe}{info.filename}"
            inventaire.append(f"{nom}\t{info.file_size} octets")
            donnees = z.read(info)
            if info.filename.lower().endswith(".zip"):
                parcourir_zip(donnees, nom + "/", inventaire)
                continue
            try:
                texte = traiter(info.filename, donnees)
            except Exception as exc:  # on continue : un fichier illisible ne doit pas tout bloquer
                texte = f"[erreur d'extraction : {exc!r}]"
            if texte is not None:
                (SORTIE / (_nom_sur(nom) + ".txt")).write_text(texte, encoding="utf-8")
                print(f"  extrait : {nom} -> {len(texte)} caractères")


def main() -> int:
    SORTIE.mkdir(parents=True, exist_ok=True)
    print(f"Téléchargement : {URL_ORA_ZIP}")
    rep = requests.get(URL_ORA_ZIP, headers=ENTETES, timeout=300, allow_redirects=True)
    print(f"HTTP {rep.status_code} ; type {rep.headers.get('Content-Type')} ; {len(rep.content)} octets")
    if rep.status_code != 200 or not rep.content.startswith(b"PK"):
        print("Échec : la réponse n'est pas une archive ZIP.")
        print(rep.text[:500])
        return 1
    inventaire: list[str] = []
    parcourir_zip(rep.content, "", inventaire)
    (SORTIE / "00_inventaire.txt").write_text(
        f"Source : {URL_ORA_ZIP} (ORA, CC BY 4.0)\n" + "\n".join(inventaire) + "\n", encoding="utf-8"
    )
    print("Inventaire de l'archive :")
    for ligne in inventaire:
        print("  " + ligne)
    return 0


if __name__ == "__main__":
    sys.exit(main())
