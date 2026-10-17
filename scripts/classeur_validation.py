"""Classeurs Excel de validation du gel 2 (charte §3.6) — export et ré-import.

Export : `docs/validation/validation_<dimension>.csv` → `docs/validation/classeurs/validation_<dimension>.xlsx`
    python scripts/classeur_validation.py exporter --dimension securite [--toutes]
Import : relit les colonnes `decision_leo` et `commentaire_leo` saisies dans le classeur et les
reporte dans le CSV (les autres colonnes du CSV ne sont pas touchées) :
    python scripts/classeur_validation.py importer --dimension securite [--fichier chemin.xlsx]

Principe (D20, validation par exception) : chaque ligne porte la recommandation de l'agent et la
décision qui s'appliquera par défaut ; Léo ne remplit `decision_leo` que là où il diverge.
Décision par défaut (D41, méthode 33 v1.3) : `socle_propose: oui` (recommandation `retenir`) → `retenu_socle` ;
autre `retenir` ou `retenir_variante` → `retenu` ; `ecarter` → `ecarte`. Repli pour une fiche sans
`socle_propose` (dimension dont la correction n'est pas fusionnée) : règle D37, `retenir` + variante principale +
incontournable → `retenu_socle`.
Le classeur montre le sous-thème et le socle proposé de chaque fiche (lus dans les fiches YAML si le tableau est
encore au format v1.2) ; la feuille « Synthèse » compte les `retenu_socle` par sous-thème (par défaut et après les
décisions saisies).
Conservation des décisions : l'export écrit dans le classeur les `decision_leo`/`commentaire_leo` du CSV ; si un
classeur existe déjà, les décisions qu'il porte et que le CSV n'a pas encore (classeur rempli mais pas encore
importé) sont reprises au lieu d'être écrasées ; l'import ne vide jamais une décision du CSV.
"""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RACINE = Path(__file__).resolve().parent.parent
DOSSIER_CSV = RACINE / "docs" / "validation"
DOSSIER_XLSX = DOSSIER_CSV / "classeurs"

DECISIONS = ["retenu", "retenu_socle", "ecarte", "a_discuter"]
POLICE = "Arial"

COLONNES = [
    # (clé CSV, en-tête affiché, largeur)
    ("sous_theme", "Sous-thème", 22),
    ("famille", "Famille", 22),
    ("id", "Identifiant", 28),
    ("libelle", "Libellé", 44),
    ("variante_principale", "Variante principale", 11),
    ("definition_courte", "Définition courte", 60),
    ("sens", "Sens", 8),
    ("unite", "Unité", 18),
    ("source", "Source", 36),
    ("serie_id", "Série (identifiant)", 30),
    ("periode", "Période", 22),
    ("type", "Type", 10),
    ("declaratif", "Déclaratif", 10),
    ("composite", "Composite", 10),
    ("incontournable", "Incontournable", 13),
    ("socle_propose", "Socle proposé", 11),
    ("recommandation", "Recommandation de l'agent", 16),
    ("motif", "Motif de la recommandation", 70),
    ("decision_par_defaut", "Décision par défaut (si vous ne remplissez rien)", 18),
    ("decision_leo", "Votre décision (liste déroulante)", 20),
    ("commentaire_leo", "Votre commentaire", 40),
]


def decision_par_defaut(ligne: dict) -> str:
    rec = (ligne.get("recommandation") or "").strip()
    principale = ligne.get("variante_principale") == "oui"
    if rec == "retenir":
        socle = (ligne.get("socle_propose") or "").strip()
        if socle:  # D41 : retenu_socle = socle_propose: oui (le vérificateur le réserve aux principales)
            return "retenu_socle" if socle == "oui" else "retenu"
        if principale and ligne.get("incontournable") == "oui":  # repli D37 (fiche sans socle_propose)
            return "retenu_socle"
        return "retenu"
    if rec == "retenir_variante":
        return "retenu"
    if rec == "ecarter":
        return "ecarte"
    return "a_discuter"


def completer_depuis_fiches(dimension: str, lignes: list[dict]) -> list[dict]:
    """Tableau antérieur à la méthode v1.3 (sans `sous_theme` ni `socle_propose`) : lit ces champs dans les fiches."""
    if not lignes or ("sous_theme" in lignes[0] and "socle_propose" in lignes[0]):
        return lignes
    import yaml  # noqa: WPS433

    fiches = {}
    for chemin in (RACINE / "indicateurs" / dimension).glob("*.yaml"):
        if not chemin.name.startswith("_"):
            f = yaml.safe_load(chemin.read_text(encoding="utf-8")) or {}
            fiches[f.get("id", chemin.stem)] = f
    for l in lignes:
        f = fiches.get(l.get("id"), {})
        l.setdefault("sous_theme", f.get("sous_theme", ""))
        l.setdefault("socle_propose", f.get("socle_propose", ""))
    return lignes


def lire_csv(dimension: str) -> list[dict]:
    chemin = DOSSIER_CSV / f"validation_{dimension}.csv"
    with chemin.open(encoding="utf-8", newline="") as f:
        return completer_depuis_fiches(dimension, list(csv.DictReader(f)))


def trier(lignes: list[dict]) -> list[dict]:
    # par sous-thème, puis famille, variante principale d'abord, puis identifiant
    return sorted(lignes, key=lambda l: (l.get("sous_theme", ""), l.get("famille", ""),
                                         0 if l.get("variante_principale") == "oui" else 1, l.get("id", "")))


def sous_themes(lignes: list[dict]) -> list[str]:
    """Sous-thèmes distincts, dans l'ordre du classeur."""
    vus: list[str] = []
    for l in lignes:
        st = l.get("sous_theme", "") or ""
        if st not in vus:
            vus.append(st)
    return vus


def lire_decisions_classeur(fichier: Path) -> dict[str, tuple[str, str]]:
    """Décisions et commentaires saisis dans un classeur (onglet « Validation »), par identifiant ; une décision
    hors liste est ignorée (avertissement)."""
    wb = load_workbook(fichier, data_only=True)
    ws = wb["Validation"]
    entetes = {str(ws.cell(row=1, column=j).value): j for j in range(1, ws.max_column + 1)}
    j_id = entetes["Identifiant"]
    j_dec = entetes["Votre décision (liste déroulante)"]
    j_com = entetes["Votre commentaire"]
    decisions = {}
    for i in range(2, ws.max_row + 1):
        ident = ws.cell(row=i, column=j_id).value
        if not ident:
            continue
        dec = str(ws.cell(row=i, column=j_dec).value or "").strip()
        com = str(ws.cell(row=i, column=j_com).value or "").strip()
        if dec and dec not in DECISIONS:
            print(f"avertissement — {ident} : décision « {dec} » non admise, ignorée", file=sys.stderr)
            dec = ""
        decisions[str(ident).strip()] = (dec, com)
    return decisions


def reprendre_decisions_classeur(lignes: list[dict], fichier: Path) -> int:
    """Avant de réécrire un classeur existant : reprend les décisions qu'il porte et que le CSV n'a pas (classeur
    rempli mais pas encore importé). En cas de désaccord, la valeur du CSV est gardée et signalée. Renvoie le nombre
    de lignes reprises."""
    if not fichier.exists():
        return 0
    try:
        existantes = lire_decisions_classeur(fichier)
    except (KeyError, ValueError, OSError) as exc:  # classeur illisible ou d'un autre format
        print(f"avertissement — {fichier.name} existant illisible ({exc}) : aucune décision reprise", file=sys.stderr)
        return 0
    n = 0
    ids = set()
    for l in lignes:
        ids.add(l.get("id"))
        dec, com = existantes.get(l.get("id"), ("", ""))
        if not (dec or com):
            continue
        actuel = ((l.get("decision_leo") or "").strip(), (l.get("commentaire_leo") or "").strip())
        if not any(actuel):
            l["decision_leo"], l["commentaire_leo"] = dec, com
            print(f"avertissement — {l.get('id')} : décision du classeur existant non encore importée, conservée "
                  f"({dec or '—'})", file=sys.stderr)
            n += 1
        elif actuel != (dec, com):
            print(f"avertissement — {l.get('id')} : le classeur existant ({dec}) diffère du CSV ({actuel[0]}) : "
                  "valeur du CSV gardée", file=sys.stderr)
    for ident, (dec, com) in existantes.items():
        if ident not in ids and (dec or com):
            print(f"avertissement — {ident} : décision du classeur existant sans ligne dans le CSV", file=sys.stderr)
    return n


def exporter(dimension: str, libelle_dimension: str | None = None) -> Path:
    lignes = trier(lire_csv(dimension))
    DOSSIER_XLSX.mkdir(parents=True, exist_ok=True)
    sortie = DOSSIER_XLSX / f"validation_{dimension}.xlsx"
    reprendre_decisions_classeur(lignes, sortie)

    wb = Workbook()
    # --- Feuille 1 : mode d'emploi -------------------------------------------------
    me = wb.active
    me.title = "Mode d'emploi"
    titre = libelle_dimension or dimension
    texte = [
        (f"Gel 2 — validation des indicateurs de la dimension « {titre} »", True),
        ("", False),
        ("Comment remplir ce classeur (validation par exception, charte §3.6, décision D20)", True),
        ("1. Ouvrez l'onglet « Validation ». Chaque ligne est un indicateur candidat (une fiche YAML dans le dépôt), regroupé par sous-thème puis par famille ; la variante principale de chaque famille est en tête.", False),
        ("2. Les colonnes grises sont informatives. La colonne « Recommandation de l'agent » et son motif résument l'analyse ; la colonne « Décision par défaut » indique ce qui s'appliquera si vous ne remplissez rien.", False),
        ("   « Socle proposé » = oui : l'agent propose de cocher l'indicateur par défaut (décision D41 : les incontournables et une fiche principale par sous-thème de la dimension, 8 au plus) ; sa décision par défaut est alors retenu_socle. L'onglet « Synthèse » compte le socle par sous-thème.", False),
        ("3. Ne remplissez la colonne jaune « Votre décision » que là où vous divergez ou hésitez. Valeurs possibles (liste déroulante) :", False),
        ("     retenu = disponible dans l'outil, non coché par défaut ;", False),
        ("     retenu_socle = disponible ET coché par défaut dans la configuration de référence ;", False),
        ("     ecarte = non retenu (indiquez pourquoi dans le commentaire) ;", False),
        ("     a_discuter = vous voulez en parler avant de trancher (la session mère vous posera la question).", False),
        ("4. Le commentaire est libre ; il est lu et conservé dans la fiche.", False),
        ("5. Enregistrez le fichier au même endroit, sous le même nom : la session mère le relit automatiquement.", False),
        ("", False),
        ("Rappels", True),
        ("— Une famille regroupe des variantes d'un même concept (autre seuil, autre source, ventilation) : en général une seule variante est cochée par défaut.", False),
        ("— Le socle est une proposition : pour l'élargir, mettez retenu_socle sur d'autres lignes ; pour le réduire, mettez retenu sur une ligne proposée.", False),
        ("— « Incontournable » = présent dans au moins trois cadres de référence (Doughnut, Stiglitz-Sen-Fitoussi, OCDE, Eurostat, INSEE).", False),
        ("— « Déclaratif » = donnée d'enquête d'opinion ou de ressenti ; « Composite » = indice agrégé d'un producteur tiers (V-Dem, CPI…).", False),
        ("— Sens « ambigu » : l'agent n'a pas pu établir si plus haut est mieux ; ces lignes sont écartées par défaut.", False),
        ("— Les bornes (seuil minimal, objectif) ne sont pas traitées ici : elles viendront après le gel 2 (socle v1).", False),
        ("", False),
        ("Exemple de ligne remplie : décision « ecarte », commentaire « doublon de la série INSEE équivalente ; garder la variante INSEE ».", False),
    ]
    for i, (t, gras) in enumerate(texte, 1):
        c = me.cell(row=i, column=1, value=t)
        c.font = Font(name=POLICE, bold=gras, size=12 if i == 1 else 10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    me.column_dimensions["A"].width = 140

    # --- Feuille 2 : validation ----------------------------------------------------
    ws = wb.create_sheet("Validation")
    entete_fill = PatternFill("solid", fgColor="1F3864")
    info_fill = PatternFill("solid", fgColor="F2F2F2")
    saisie_fill = PatternFill("solid", fgColor="FFF2CC")
    bord = Border(bottom=Side(style="thin", color="BFBFBF"))
    for j, (_, en_tete, largeur) in enumerate(COLONNES, 1):
        c = ws.cell(row=1, column=j, value=en_tete)
        c.font = Font(name=POLICE, bold=True, color="FFFFFF")
        c.fill = entete_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = largeur
    ws.row_dimensions[1].height = 42

    col = {cle: j for j, (cle, _, _) in enumerate(COLONNES, 1)}
    for i, l in enumerate(lignes, 2):
        l = dict(l)
        l["decision_par_defaut"] = decision_par_defaut(l)
        for cle, j in col.items():
            c = ws.cell(row=i, column=j, value=l.get(cle, "") or "")
            c.font = Font(name=POLICE, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = bord
            c.fill = saisie_fill if cle in ("decision_leo", "commentaire_leo") else info_fill
        ws.row_dimensions[i].height = 60
        # ligne de variante principale : identifiant en gras
        if l.get("variante_principale") == "oui":
            ws.cell(row=i, column=col["id"]).font = Font(name=POLICE, size=10, bold=True)

    n = len(lignes) + 1
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True, showErrorMessage=True,
                        errorTitle="Valeur non admise", error="Choisissez une valeur dans la liste : " + ", ".join(DECISIONS))
    ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(col['decision_leo'])}2:{get_column_letter(col['decision_leo'])}{max(n, 2)}")
    ws.freeze_panes = ws.cell(row=2, column=col["libelle"] + 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLONNES))}{max(n, 2)}"
    # mise en évidence des lignes où une décision a été saisie
    lettre_dec = get_column_letter(col["decision_leo"])
    ws.conditional_formatting.add(
        f"A2:{get_column_letter(len(COLONNES))}{max(n, 2)}",
        FormulaRule(formula=[f"${lettre_dec}2<>\"\""], fill=PatternFill("solid", fgColor="E2EFDA")),
    )
    ws.cell(row=1, column=col["decision_leo"]).comment = Comment("À remplir seulement si vous divergez de la décision par défaut. Liste : retenu, retenu_socle, ecarte, a_discuter.", "session mère")

    # --- Feuille 3 : synthèse (formules) -------------------------------------------
    sy = wb.create_sheet("Synthèse")

    def plage(cle: str) -> str:
        lettre = get_column_letter(col[cle])
        return f"Validation!${lettre}$2:${lettre}${n}"

    plage_rec, plage_def, plage_dec = plage("recommandation"), plage("decision_par_defaut"), plage("decision_leo")
    plage_fam, plage_st, plage_id = plage("famille"), plage("sous_theme"), plage("id")
    lignes_syn = [
        ("Dimension", titre),
        ("Indicateurs candidats (fiches)", f"=COUNTA({plage_id})"),
        ("Familles distinctes", f"=SUMPRODUCT(({plage_fam}<>\"\")/COUNTIF({plage_fam},{plage_fam}&\"\"))"),
        ("", ""),
        ("Recommandations de l'agent", ""),
        ("  retenir", f"=COUNTIF({plage_rec},\"retenir\")"),
        ("  retenir_variante", f"=COUNTIF({plage_rec},\"retenir_variante\")"),
        ("  ecarter", f"=COUNTIF({plage_rec},\"ecarter\")"),
        ("", ""),
        ("Socle proposé par l'agent (D41)", f"=COUNTIF({plage('socle_propose')},\"oui\")"),
        ("", ""),
        ("Décisions par défaut", ""),
        ("  retenu_socle", f"=COUNTIF({plage_def},\"retenu_socle\")"),
        ("  retenu", f"=COUNTIF({plage_def},\"retenu\")"),
        ("  ecarte", f"=COUNTIF({plage_def},\"ecarte\")"),
        ("  a_discuter", f"=COUNTIF({plage_def},\"a_discuter\")"),
        ("", ""),
        ("Vos décisions saisies (exceptions)", f"=COUNTA({plage_dec})"),
        ("  dont retenu_socle", f"=COUNTIF({plage_dec},\"retenu_socle\")"),
        ("  dont retenu", f"=COUNTIF({plage_dec},\"retenu\")"),
        ("  dont ecarte", f"=COUNTIF({plage_dec},\"ecarte\")"),
        ("  dont a_discuter", f"=COUNTIF({plage_dec},\"a_discuter\")"),
    ]
    for i, (a, b) in enumerate(lignes_syn, 1):
        ca = sy.cell(row=i, column=1, value=a)
        cb = sy.cell(row=i, column=2, value=b)
        ca.font = Font(name=POLICE, bold=not a.startswith("  ") and a != "", size=10)
        cb.font = Font(name=POLICE, size=10)

    # socle par sous-thème (D41) : par défaut, et après les décisions saisies (la décision saisie l'emporte)
    debut = len(lignes_syn) + 2
    entetes_st = ["Socle par sous-thème (D41)", "Fiches", "retenu_socle par défaut", "retenu_socle après vos décisions"]
    for j, t in enumerate(entetes_st, 1):
        sy.cell(row=debut, column=j, value=t).font = Font(name=POLICE, bold=True, size=10)
    ligne = debut
    for ligne, st in enumerate(sous_themes(lignes), debut + 1):
        ref = f"$A{ligne}"
        valeurs = [st or "(sans sous-thème)",
                   f"=COUNTIF({plage_st},{ref})",
                   f"=COUNTIFS({plage_st},{ref},{plage_def},\"retenu_socle\")",
                   f"=SUMPRODUCT(({plage_st}={ref})*((({plage_dec}=\"\")*({plage_def}=\"retenu_socle\"))"
                   f"+({plage_dec}=\"retenu_socle\")))"]
        for j, v in enumerate(valeurs, 1):
            sy.cell(row=ligne, column=j, value=v).font = Font(name=POLICE, size=10)
    total = ligne + 1
    sy.cell(row=total, column=1, value="Total").font = Font(name=POLICE, bold=True, size=10)
    for j in (2, 3, 4):
        lettre = get_column_letter(j)
        sy.cell(row=total, column=j, value=f"=SUM({lettre}{debut + 1}:{lettre}{ligne})").font = Font(name=POLICE, bold=True, size=10)
    sy.column_dimensions["A"].width = 40
    sy.column_dimensions["B"].width = 18
    sy.column_dimensions["C"].width = 22
    sy.column_dimensions["D"].width = 30
    sy.cell(row=total + 2, column=1, value="Les comptes se recalculent à l'ouverture dans Excel.").font = Font(name=POLICE, italic=True, size=9)

    wb.save(sortie)
    return sortie


def importer(dimension: str, fichier: Path | None = None) -> int:
    chemin_csv = DOSSIER_CSV / f"validation_{dimension}.csv"
    fichier = fichier or (DOSSIER_XLSX / f"validation_{dimension}.xlsx")
    decisions = lire_decisions_classeur(fichier)
    with chemin_csv.open(encoding="utf-8", newline="") as f:
        lecteur = csv.DictReader(f)
        champs = lecteur.fieldnames or []
        lignes = list(lecteur)
    n = 0
    for l in lignes:
        dec, com = decisions.get(l["id"], ("", ""))
        # une cellule vide du classeur ne vide jamais une valeur déjà importée
        nouveau = (dec or l.get("decision_leo", ""), com or l.get("commentaire_leo", ""))
        if (dec or com) and nouveau != (l.get("decision_leo", ""), l.get("commentaire_leo", "")):
            l["decision_leo"], l["commentaire_leo"] = nouveau
            n += 1
    with chemin_csv.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=champs, lineterminator="\n")
        w.writeheader()
        w.writerows(lignes)
    print(f"{dimension} : {n} décision(s) ou commentaire(s) reportés dans {chemin_csv.name}")
    return n


def libelles_dimensions() -> dict[str, str]:
    try:
        import yaml  # noqa: WPS433

        d = yaml.safe_load((RACINE / "indicateurs" / "dimensions.yaml").read_text(encoding="utf-8"))
        return {x["id"]: x["libelle"] for x in d["dimensions"]}
    except Exception:  # noqa: BLE001
        return {}


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sp = p.add_subparsers(dest="commande", required=True)
    e = sp.add_parser("exporter")
    e.add_argument("--dimension")
    e.add_argument("--toutes", action="store_true", help="toutes les dimensions ayant un CSV de validation")
    i = sp.add_parser("importer")
    i.add_argument("--dimension", required=True)
    i.add_argument("--fichier")
    args = p.parse_args()
    lib = libelles_dimensions()
    if args.commande == "exporter":
        if args.toutes:
            dims = sorted(c.stem.replace("validation_", "") for c in DOSSIER_CSV.glob("validation_*.csv"))
        elif args.dimension:
            dims = [args.dimension]
        else:
            p.error("--dimension ou --toutes")
        for d in dims:
            print("écrit", exporter(d, lib.get(d)))
    else:
        importer(args.dimension, Path(args.fichier) if args.fichier else None)
    return 0


if __name__ == "__main__":
    sys.exit(main())
